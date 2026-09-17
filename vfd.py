from search import collection
import openpyxl

wb1 = openpyxl.load_workbook(r"D:\FactoryKA\documents\VFD3000_Fault_Alarm_Codes.xlsx")
wb2 = openpyxl.load_workbook(r"D:\FactoryKA\documents\VFD_Fault_Alarm_Code_Table.xlsx")

print("=== VFD3000_Fault_Alarm_Codes.xlsx ===")
for row in wb1.active.iter_rows(min_row=1, max_row=5, values_only=True):
    print(row)

print("\n=== VFD_Fault_Alarm_Code_Table.xlsx ===")
for row in wb2.active.iter_rows(min_row=1, max_row=5, values_only=True):
    print(row)