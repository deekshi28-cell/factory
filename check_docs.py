import fitz  # pymupdf
import os
import docx
import openpyxl

folder = r"D:\FactoryKA\documents"
total_pages = 0

print(f"{'File':<50}{'Pages':<10}{'Type':<12}{'Text?':<10}")

for fname in os.listdir(folder):
    path = os.path.join(folder, fname)
    lower = fname.lower()

    if lower.endswith(".pdf"):
        doc = fitz.open(path)
        pages = len(doc)
        total_pages += pages
        has_text = len(doc[0].get_text().strip()) > 20
        print(f"{fname[:48]:<50}{pages:<10}{'PDF':<12}{'Yes' if has_text else 'SCANNED?':<10}")
        doc.close()

    elif lower.endswith(".docx"):
        d = docx.Document(path)
        para_count = len(d.paragraphs)
        table_count = len(d.tables)
        print(f"{fname[:48]:<50}{'N/A':<10}{'DOCX':<12}{f'{para_count} para, {table_count} tbl':<10}")

    elif lower.endswith(".xlsx"):
        wb = openpyxl.load_workbook(path, data_only=True)
        sheet_count = len(wb.sheetnames)
        print(f"{fname[:48]:<50}{'N/A':<10}{'XLSX':<12}{f'{sheet_count} sheet(s)':<10}")

print(f"\nTotal PDF pages: {total_pages}")
print(f"Note: page count only applies to PDFs. Word/Excel shown separately above.")